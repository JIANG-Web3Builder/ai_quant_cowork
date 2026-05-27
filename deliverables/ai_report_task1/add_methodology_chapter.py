from __future__ import annotations

from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


REPORT_DIR = Path(__file__).resolve().parent / "outputs_enhanced" / "report"
CHAPTER_TITLE = "十、关键步骤汇总与AI执行方法论"
COLLAB_TITLE = "（六）AI协同搭建研究工程底座"
FONT = "Microsoft YaHei"
BLUE = RGBColor(31, 78, 121)


def find_report() -> Path:
    candidates = [
        path
        for path in REPORT_DIR.glob("*.docx")
        if "ai" in path.name.lower()
        and "before-methodology" not in path.name.lower()
        and not path.name.startswith("~$")
    ]
    if not candidates:
        raise FileNotFoundError(f"No renamed AI report found in {REPORT_DIR}")
    return max(candidates, key=lambda path: path.stat().st_mtime)


def set_font(run, size: float = 10.5, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_font(run, 14 if level == 1 else 12, True, BLUE)


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(21)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_font(run, 10.5)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def style_cell(cell, bold: bool = False, size: float = 8.5, align_center: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if align_center else WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            set_font(run, size, bold)


def add_table(document: Document, title: str, headers: list[str], rows: list[list[str]]) -> None:
    caption = document.add_paragraph()
    run = caption.add_run(title)
    set_font(run, 10, True)
    table = document.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
        shade_cell(table.rows[0].cells[idx], "D9EAF7")
        style_cell(table.rows[0].cells[idx], bold=True, align_center=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            style_cell(cells[idx], size=8.0)
    table.autofit = True


def append_methodology_chapter(document: Document) -> None:
    document.add_page_break()
    add_heading(document, CHAPTER_TITLE, 1)
    add_body(
        document,
        "本章对前述选股与择时AI落地流程进行归纳，形成从任务边界设定、数据读取、因子构建、滚动验证选模、结果复盘到报告输出的闭环方法。该闭环强调两个原则：一是所有模型输入均来自可复现的数据底座，二是AI只承担归纳、筛选、解释和迭代增强功能，不替代数据校验、样本外验证和研究结论审定。",
    )

    add_heading(document, "（一）关键步骤汇总", 2)
    add_body(
        document,
        "本方案的关键步骤可以拆分为六个环节。每个环节均包含明确输入、AI执行动作、输出结果和校验口径，使研究过程从一次性文本讨论转化为可复盘的标准化流程。",
    )
    add_table(
        document,
        "表12：关键步骤、AI执行动作与输出结果汇总",
        ["环节", "输入材料", "AI执行动作", "输出结果", "校验口径"],
        [
            ["任务边界设定", "研究目标、数据范围、报告用途、可用数据目录", "归纳任务目标，拆分选股、择时、复盘和报告四类工作边界", "形成研究边界、样本区间、训练区间和样本外区间", "确认不使用样本外数据制定权重和参数"],
            ["数据底座梳理", "日频行情、财务指标、交易约束、指数数据、60分钟价量数据", "识别可用字段、数据覆盖范围和潜在口径风险", "形成数据覆盖表和未来函数检查表", "检查财务公告日、滚动窗口和分钟样本选择口径"],
            ["候选因子构建", "估值、质量、成长、资本结构、趋势反转、波动、流动性、价格位置和分钟结构字段", "将原始字段转化为标准化因子，并按横截面做去极值和标准化", "形成候选因子池、因子分组摘要和IC结果", "检查因子仅使用交易日当日及以前信息"],
            ["AI辅助筛选", "训练期IC、IC_IR、正IC占比、因子相关性、分组稳定性", "把因子归类为核心加权、辅助过滤、观察保留、反向验证或降权", "形成因子动作清单、权重表和分组稳定性表", "权重仅由训练期表现生成"],
            ["模型组合验证", "训练期候选标签、TopN持仓数量、交易成本、基准收益", "枚举多因子IC加权与Ridge横截面模型，并按滚动验证目标函数排序", "选定稳定性更好的候选方案并做样本外验证", "样本外只作为验证，不参与模型选择"],
            ["报告复盘输出", "回测摘要、分层收益、净值曲线、市场状态、分钟特征", "将结果转化为管理层可读的结论、表格和图形", "形成正式方案报告和可复现输出目录", "检查表格、图形、成本口径和章节表述一致"],
        ],
    )

    add_heading(document, "（二）AI执行对话框架", 2)
    add_body(
        document,
        "AI执行并非一次性生成结论，而是围绕“任务澄清—数据读取—假设生成—实验执行—结果解释—迭代优化”的连续对话框架展开。对话内容在报告中不以聊天记录形式展示，而是沉淀为可执行指令、结构化输出和复盘结论。",
    )
    add_table(
        document,
        "表13：AI执行对话与结构化输出示例",
        ["阶段", "对话指令口径", "AI输出要求", "进入下一步的条件"],
        [
            ["任务澄清", "围绕选股与择时中的AI可应用环节，限定数据目录、报告目标和交付形态", "输出任务边界、数据范围、研究模块和报告结构", "边界清晰，训练期和样本外区间明确"],
            ["数据检查", "读取日频数据、财务数据、交易约束和60分钟价量数据，识别可用字段", "输出数据覆盖、字段口径、缺失情况和潜在风险", "字段能支撑因子构建，且无明显未来函数"],
            ["因子假设", "基于强势样本、历史有效因子和数据字段生成候选因子", "输出因子名称、经济含义、计算方式和预期方向", "因子可计算、方向可解释、口径可复现"],
            ["实验执行", "计算IC、分层收益、因子相关性和组合回测", "输出CSV、图形、摘要指标和异常点", "训练期和样本外结果分开呈现"],
            ["复盘解释", "根据收益、回撤、换手、风格和市场状态解释模型表现", "输出优势、短板、适用环境和失效情景", "结论与数据证据对应，不使用空泛判断"],
            ["报告固化", "把实验过程写成方案报告章节，保留方法论和关键数据表", "输出正式Word报告、表格和图形附件", "报告用语为方案正文口吻，指标和成本口径一致"],
        ],
    )

    add_heading(document, "（三）具体执行路径", 2)
    add_body(
        document,
        "具体执行路径采用“先数据、后因子、再模型、最后复盘”的顺序。第一步读取研究数据库和分钟数据，完成字段识别和样本筛选；第二步构建候选因子并进行横截面标准化；第三步用训练期IC和稳定性生成因子权重或机器学习系数；第四步在训练期内按滚动窗口枚举不同目标期限、模型类型和持仓数量；第五步将选定模型放入最终样本外区间验证；第六步把样本外表现、分层曲线、分钟特征和择时状态统一写入报告。",
    )
    add_body(
        document,
        "执行过程中，AI的作用集中在信息压缩和结构化表达。对于因子研究，AI负责把大量IC、正IC占比和相关性结果压缩为动作清单；对于模型复盘，AI负责把收益、回撤、换手和状态标签转化为可读解释；对于报告生成，AI负责把数据表和图形组织成正式方案文本。",
    )

    add_heading(document, "（四）优化前后对比", 2)
    add_body(
        document,
        "本轮优化经历了从初版等权复合因子到滚动验证驱动模型选择的迭代。初版模型以固定权重组合为主，样本外表现更容易受因子方向和持仓结构影响；优化版引入技术因子扩展、机器学习截面模型、训练期内部滚动验证和更严格的未来函数隔离，组合稳定性和报告可解释性均有提升。",
    )
    add_table(
        document,
        "表14：优化前后关键结果对比",
        ["项目", "初版数据支撑", "增强优化版", "变化说明"],
        [
            ["模型形成方式", "人工设定复合因子权重", "训练期IC加权与Ridge横截面模型并行评估", "从经验权重转向可审计的模型比较"],
            ["训练/验证隔离", "已区分研究期和验证结果，但标签边界较粗", "模型权重和参数仅使用训练期，并在训练期内部做滚动验证", "样本外验证口径更严格"],
            ["候选因子范围", "以基础因子为主", "加入技术、流动性、价格位置和分钟结构等扩展特征", "因子覆盖更完整"],
            ["模型筛选方式", "主要依据单段训练结果", "优先依据滚动验证稳定性，再看全训练期复核", "降低单段样本筛优偏差"],
            ["结果表达", "以实证补充为主", "同步展示分组稳定性、模型选择、回测和方法论", "更接近正式方案报告"],
            ["报告表达", "以实证补充为主", "补充关键步骤汇总、执行框架和方法论", "更接近正式方案报告"],
        ],
    )

    add_heading(document, "（五）方法论沉淀", 2)
    add_body(
        document,
        "本方案沉淀的方法论可以概括为“三层闭环”。第一层是数据闭环，即所有研究结论均对应可追溯的数据表、图形和计算脚本；第二层是模型闭环，即因子权重、模型参数和样本外验证严格分离；第三层是复盘闭环，即每一次模型输出均转化为可解释的问题清单、适用场景和下一轮迭代方向。",
    )
    add_body(
        document,
        "在该方法论下，AI不直接输出不可验证的投资判断，而是把研究过程中的非结构化讨论转化为结构化实验，把零散指标转化为可读结论，把单次回测转化为可持续迭代的研究流程。该定位能够同时满足研究效率、模型复盘、因子提升和管理层汇报四类需求。",
    )

    append_collaboration_section(document)


def append_collaboration_section(document: Document) -> None:
    add_heading(document, COLLAB_TITLE, 2)
    add_body(
        document,
        "本项目的实施过程本身也体现了AI在研究流程优化中的应用价值。研究数据底座、代码工程结构、数据读取脚本、因子计算脚本、回测验证模块、图表生成模块和Word报告生成流程，均在人机协同方式下逐步搭建完成。研究人员负责确定任务目标、数据边界、业务含义和结果取舍，AI负责把需求拆解为可执行代码、自动读取本地数据、生成中间表和图形，并将计算结果组织为方案报告。",
    )
    add_body(
        document,
        "该协作方式使任务从单纯撰写文字方案，扩展为“数据工程—因子研究—滚动验证—结果复盘—报告输出”的一体化工作流。AI在其中承担了工程搭建和研究辅助双重角色：一方面快速形成可复用的本地研究框架，降低数据整理和重复计算成本；另一方面把模型测试结果转化为可解释结论，提升报告的证据密度和复盘效率。",
    )
    add_table(
        document,
        "表15：本项目中AI协同完成的工作内容",
        ["工作模块", "人工输入", "AI协同执行内容", "形成结果"],
        [
            ["研究工程搭建", "明确本地数据目录、研究目标和交付报告形态", "搭建脚本目录、输出目录和可复现运行流程", "形成面向选股与择时研究的数据支撑工程"],
            ["数据接入与检查", "指定60分钟数据、Tushare数据和已有研究库", "读取DuckDB、Parquet和Word文件，识别字段和覆盖范围", "形成数据覆盖表和口径检查结果"],
            ["因子计算与验证", "确定可用因子类型和训练/样本外边界", "编写因子计算、IC检验、分层收益和组合回测逻辑", "形成因子IC表、权重表、回测摘要和图形"],
            ["模型复盘与优化", "判断模型表现是否需要优化和防未来函数检查", "枚举目标期限和持仓数量，比较优化前后结果", "形成训练期模型选择和样本外验证结果"],
            ["报告生成", "确定报告应呈现为正式方案而非对话记录", "将数据表、图形和方法论写入Word", "形成可直接用于汇报的方案文档"],
        ],
    )


def main() -> None:
    report_path = find_report()
    backup_path = report_path.with_name(report_path.stem + "-before-methodology-addition.docx")
    if not backup_path.exists():
        copy2(report_path, backup_path)

    document = Document(report_path)
    text = "\n".join(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
    if CHAPTER_TITLE not in text:
        append_methodology_chapter(document)
        document.save(report_path)

    print(f"report={report_path}")
    print(f"backup={backup_path}")
    print(f"chapter_present={CHAPTER_TITLE in '\\n'.join(p.text for p in Document(report_path).paragraphs)}")


if __name__ == "__main__":
    main()
